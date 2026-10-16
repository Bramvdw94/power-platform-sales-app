"""Bouwt presentatie-sales-app.pptx in de Fellowmind-huisstijl.

Gebruik (vanuit de map presentation/):
    pip install python-pptx
    python maak-pptx.py

Inhoud komt uit presentatie-sales-app.md; cijfers zijn gecontroleerd tegen GitHub Issues.
De diagrammen worden als native PowerPoint-vormen getekend (lijnstijl), niet als afbeelding.
Installeer Montserrat en Roboto (Google Fonts) voordat je de .pptx opent of presenteert.
"""
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.parts.image import Image as PptxImage
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
IMG = HERE / "afbeeldingen"
VIDEO = HERE / "video" / "demo-sales-app.mp4"
LOGO_SVG = HERE / "huisstijl" / "fellowmind-logo.svg"
OUT = HERE / "presentatie-sales-app.pptx"
REPO_URL = "https://github.com/Bramvdw94/power-platform-sales-app"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

# Merkkleuren (Brand Guidelines v3.0)
WHITE, LIGHT, DARK, BLACK = "FFFFFF", "F8F8F8", "1F1F1F", "000000"
GREY = "5F5F5F"   # secundaire tekst (uitbreiding: tint van zwart)
RULE = "D6D6D6"   # haarlijnen tussen items
ACCENT = "C514B4"  # Fellow-highlight, voor labels en de ene nadruk per dia
FELLOW = ["F05027", "FF455B", "D53B91", "C514B4"]

# Fonts: Montserrat voor koppen, Roboto voor tekst
HEAD, HEAD_SB, HEAD_MD = "Montserrat", "Montserrat SemiBold", "Montserrat Medium"
BODY, BODY_MD = "Roboto", "Roboto Medium"

# 16:9-grid: 9 kolommen, marge = kolom/3, goot = marge/2, radius = marge/2
SW, SH = 13.333, 7.5
COL = SW / 11
M = COL / 3
G = M / 2
R = M / 2
TOP = 1.95                      # bovenkant van de inhoud
BOTTOM = SH - M - 0.45          # onderkant van de inhoud (boven de voettekst)


def X(i):
    return M + i * (COL + G)


def SPAN(n):
    return n * COL + (n - 1) * G


RIGHT = X(0) + SPAN(9)


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- basisvormen

def strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def add_text(slide, x, y, w, h, paras, *, font=BODY, size=16, color=DARK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.12, space_after=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, (str, dict)):
        paras = [paras]
    for i, p in enumerate(paras):
        if isinstance(p, str):
            p = {"text": p}
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", align)
        para.line_spacing = p.get("line_spacing", line_spacing)
        if p.get("space_after", space_after):
            para.space_after = Pt(p.get("space_after", space_after))
        if p.get("space_before"):
            para.space_before = Pt(p["space_before"])
        runs = p.get("runs") or [(p.get("text", ""), {})]
        for text, st in runs:
            r = para.add_run()
            r.text = text
            f = r.font
            f.name = st.get("font", p.get("font", font))
            f.size = Pt(st.get("size", p.get("size", size)))
            f.bold = st.get("bold", p.get("bold", bold))
            f.color.rgb = rgb(st.get("color", p.get("color", color)))
            r._r.get_or_add_rPr().set("lang", "nl-NL")
            if st.get("spc", p.get("spc")):
                r._r.get_or_add_rPr().set("spc", str(st.get("spc", p.get("spc"))))
            if st.get("link"):
                r.hyperlink.address = st["link"]
    return tb


def line(slide, x1, y1, x2, y2, color=DARK, width=0.75, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    strip_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if arrow:
        ln = c.line._get_or_add_ln()
        te = etree.SubElement(ln, qn("a:tailEnd"))
        te.set("type", "triangle")
        te.set("w", "med")
        te.set("len", "med")
    return c


def arrow(slide, x1, y1, x2, y2):
    return line(slide, x1, y1, x2, y2, color=GREY, width=1.0, arrow=True)


def box(slide, x, y, w, h, outline=DARK, fill=None, width=1.0, rounded=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    strip_style(shp)
    if rounded:
        shp.adjustments[0] = min(0.5, R / min(w, h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if outline:
        shp.line.color.rgb = rgb(outline)
        shp.line.width = Pt(width)
    else:
        shp.line.fill.background()
    return shp


def gradient_xml(angle_deg):
    stops = "".join(
        f'<a:gs pos="{round(i * 100000 / (len(FELLOW) - 1))}"><a:srgbClr val="{c}"/></a:gs>'
        for i, c in enumerate(FELLOW))
    return etree.fromstring(
        f'<a:gradFill xmlns:a="{A_NS}" rotWithShape="1"><a:gsLst>{stops}</a:gsLst>'
        f'<a:lin ang="{int(angle_deg * 60000)}" scaled="0"/></a:gradFill>')


def apply_gradient(shape, angle_deg=0):
    """Lineaire Fellow-gradient (0 = links naar rechts, 90 = boven naar onder)."""
    shape.fill.solid()
    sp_pr = shape._element.spPr
    sf = sp_pr.find(qn("a:solidFill"))
    sf.addprevious(gradient_xml(angle_deg))
    sp_pr.remove(sf)


def gradient_rect(slide, x, y, w, h, angle_deg=0):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    strip_style(shp)
    shp.line.fill.background()
    apply_gradient(shp, angle_deg)  # gradiënten altijd met scherpe hoeken
    return shp


# ---------------------------------------------------- vectorartwork (logo, slant)

def parse_svg_path(d):
    tokens = re.findall(r"[MLCHVZ]|-?\d*\.?\d+", d)
    cmds, i, cur, start, cmd = [], 0, (0.0, 0.0), (0.0, 0.0), None

    def num():
        nonlocal i
        v = float(tokens[i])
        i += 1
        return v

    while i < len(tokens):
        if tokens[i].isalpha():
            cmd = tokens[i]
            i += 1
        if cmd == "Z":
            cmds.append(("Z",))
            cur = start
        elif cmd == "M":
            cur = (num(), num())
            start = cur
            cmds.append(("M", cur))
            cmd = "L"
        elif cmd == "L":
            cur = (num(), num())
            cmds.append(("L", cur))
        elif cmd == "H":
            cur = (num(), cur[1])
            cmds.append(("L", cur))
        elif cmd == "V":
            cur = (cur[0], num())
            cmds.append(("L", cur))
        elif cmd == "C":
            p1, p2, cur = (num(), num()), (num(), num()), (num(), num())
            cmds.append(("C", p1, p2, cur))
        else:
            raise ValueError(f"Niet ondersteund padcommando: {cmd}")
    return cmds


def custgeom_xml(paths, vw, vh, scale=100):
    def pt(p):
        return f'<a:pt x="{round(p[0] * scale)}" y="{round(p[1] * scale)}"/>'

    body = []
    for cmds in paths:
        for c in cmds:
            if c[0] == "M":
                body.append(f"<a:moveTo>{pt(c[1])}</a:moveTo>")
            elif c[0] == "L":
                body.append(f"<a:lnTo>{pt(c[1])}</a:lnTo>")
            elif c[0] == "C":
                body.append(f"<a:cubicBezTo>{pt(c[1])}{pt(c[2])}{pt(c[3])}</a:cubicBezTo>")
            else:
                body.append("<a:close/>")
    return etree.fromstring(
        f'<a:custGeom xmlns:a="{A_NS}"><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>'
        f'<a:rect l="0" t="0" r="r" b="b"/><a:pathLst>'
        f'<a:path w="{round(vw * scale)}" h="{round(vh * scale)}">{"".join(body)}</a:path>'
        f'</a:pathLst></a:custGeom>')


def vector(slide, paths, vw, vh, x, y, w, h, color=None, gradient_angle=None, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    strip_style(shp)
    sp_pr = shp._element.spPr
    pg = sp_pr.find(qn("a:prstGeom"))
    pg.addprevious(custgeom_xml(paths, vw, vh))
    sp_pr.remove(pg)
    shp.line.fill.background()
    if gradient_angle is not None:
        apply_gradient(shp, gradient_angle)
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(color)
    if name:
        shp.name = name
    return shp


LOGO_PATHS = [parse_svg_path(d) for d in re.findall(r'\sd="([^"]+)"', LOGO_SVG.read_text(encoding="utf-8"))]
LOGO_VW, LOGO_VH = 480.0, 78.31
SLANT_PATH = [parse_svg_path("M143 0 H196 L53 269 H0 Z")]  # -28°, vaste verhouding 196:269


def logo(slide, x, y, w, color):
    """Het aangeleverde Fellowmind-woordlogo, zwart of wit, exact uit de SVG-paden."""
    vector(slide, LOGO_PATHS, LOGO_VW, LOGO_VH, x, y, w, w * LOGO_VH / LOGO_VW, color=color,
           name="Fellowmind logo")


def slant(slide, x, y, h, white=False):
    w = h * 196 / 269
    vector(slide, SLANT_PATH, 196, 269, x, y, w, h, color=WHITE if white else None,
           gradient_angle=None if white else 90, name="Fellowmind slant")
    return w


# ---------------------------------------------------------------- afbeeldingen

def round_corners(shape, w, h):
    pg = shape._element.spPr.find(qn("a:prstGeom"))
    pg.set("prst", "roundRect")
    av = pg.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(pg, qn("a:avLst"))
    for child in list(av):
        av.remove(child)
    gd = etree.SubElement(av, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", f"val {int(min(50000, R / min(w, h) * 100000))}")


def picture(slide, path, x, y, w, h, halign="center", valign="top", border=True):
    iw, ih = PptxImage.from_file(str(path)).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    px = {"left": x, "center": x + (w - pw) / 2, "right": x + w - pw}[halign]
    py = {"top": y, "middle": y + (h - ph) / 2, "bottom": y + h - ph}[valign]
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    round_corners(pic, pw, ph)
    if border:
        pic.line.color.rgb = rgb(RULE)
        pic.line.width = Pt(0.75)
    return px, py, pw, ph


def caption(slide, x, y, w, text):
    add_text(slide, x, y, w, 0.3, text, size=10, color=GREY)


# ------------------------------------------------------------ dia-opbouw

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def set_theme_fonts():
    """Zet de themafonts op Montserrat/Roboto, zodat nieuw getypte tekst ook in huisstijl staat."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    theme = prs.slide_master.part.part_related_by(RT.THEME)
    xml = theme.blob.decode("utf-8")
    xml = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*', r"\g<1>Montserrat", xml)
    xml = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*', r"\g<1>Roboto", xml)
    # Hyperlinks take the theme colour, not the run colour: keep them dark instead of default blue.
    xml = re.sub(r"<a:hlink>.*?</a:hlink>", f'<a:hlink><a:srgbClr val="{DARK}"/></a:hlink>', xml, flags=re.S)
    xml = re.sub(r"<a:folHlink>.*?</a:folHlink>", f'<a:folHlink><a:srgbClr val="{DARK}"/></a:folHlink>', xml, flags=re.S)
    if hasattr(theme, "_element"):
        theme._element = etree.fromstring(xml.encode("utf-8"))
    else:
        theme._blob = xml.encode("utf-8")


def set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def notes(slide, text, tijd=None):
    slide.notes_slide.notes_text_frame.text = text + (f"\n\nTijd: {tijd}" if tijd else "")


def footer(slide):
    n = len(prs.slides)
    add_text(slide, X(0), SH - M - 0.17, SPAN(5), 0.2,
             {"runs": [(f"{n}", {"font": HEAD_MD}), ("    Een CRM bouwen als functioneel consultant", {})]},
             size=9, color=GREY)
    lw = 1.15
    logo(slide, SW - M - lw, SH - M - lw * LOGO_VH / LOGO_VW, lw, BLACK)


def content_slide(label, title, bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    set_bg(s, bg)
    add_text(s, X(0), M + 0.05, SPAN(9), 0.25, label.upper(), font=HEAD_SB, size=11, color=ACCENT)
    add_text(s, X(0), M + 0.38, SPAN(9), 0.9, title, font=HEAD, size=30, color=BLACK, line_spacing=1.0)
    footer(s)
    return s


def lined_list(s, x, y, w, items, row_h, *, size=16, desc_size=12.5, numbered=False,
               title_font=BODY, first_rule=DARK):
    """Lijst gestructureerd met dunne lijnen in plaats van bullets of kaders."""
    for i, it in enumerate(items):
        yy = y + i * row_h
        line(s, x, yy, x + w, yy, color=first_rule if i == 0 else RULE)
        title, desc = it if isinstance(it, tuple) else (it, None)
        tx = x
        if numbered:
            add_text(s, x, yy + 0.12, 0.55, 0.4, f"{i + 1:02d}", font=HEAD, size=size, color=ACCENT)
            tx = x + 0.65
        paras = [{"text": title, "font": title_font, "size": size}]
        if desc:
            paras.append({"text": desc, "size": desc_size, "color": GREY, "space_before": 2})
        add_text(s, tx, yy + 0.12, w - (tx - x), row_h - 0.14, paras)
    return y + len(items) * row_h


def node(s, x, y, w, h, title, sub=None, highlight=False, title_size=14, sub_size=11, align=PP_ALIGN.LEFT):
    """Diagramknoop: monoline afgeronde rechthoek, één accentkleur voor nadruk."""
    box(s, x, y, w, h, outline=ACCENT if highlight else DARK, width=1.5 if highlight else 1.0)
    pad = 0.16
    paras = [{"text": title, "font": HEAD_SB, "size": title_size, "align": align}]
    if sub:
        paras.append({"text": sub, "size": sub_size, "color": GREY, "space_before": 3, "align": align})
    add_text(s, x + pad, y + pad * 0.8, w - 2 * pad, h - 1.6 * pad, paras,
             anchor=MSO_ANCHOR.MIDDLE if align == PP_ALIGN.CENTER else MSO_ANCHOR.TOP)


def poster_frame():
    """Posterframe voor de video: frame uit de video (ffmpeg) of de Home-schermafbeelding, 16:9."""
    from PIL import Image, ImageDraw
    tmp = Path(tempfile.mkdtemp())
    src = tmp / "frame.png"
    if shutil.which("ffmpeg"):
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-ss", "6", "-i", str(VIDEO),
                        "-frames:v", "1", str(src)], check=False)
    if src.exists():
        im = Image.open(src).convert("RGB")
    else:
        im = Image.open(IMG / "01-home.png").convert("RGB")
        w = im.width
        im = im.crop((0, 0, w, int(w * 9 / 16)))
    im = im.resize((1280, 720))
    d = ImageDraw.Draw(im)
    cx, cy, rr = 640, 360, 64
    d.ellipse((cx - rr, cy - rr, cx + rr, cy + rr), fill=(255, 255, 255), outline=(31, 31, 31), width=4)
    d.polygon([(cx - 18, cy - 28), (cx - 18, cy + 28), (cx + 30, cy)], fill=(31, 31, 31))
    out = tmp / "poster.png"
    im.save(out)
    return out


# ======================================================================== dia's

set_theme_fonts()

# 1. Titel ---------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
set_bg(s, LIGHT)
panel_w = X(5) + COL * 0.5
gradient_rect(s, 0, 0, panel_w, SH, 0)
logo(s, X(0), M + 0.1, 1.7, WHITE)
add_text(s, X(0), 2.15, SPAN(5), 0.3, {"text": "KENNISSESSIE POWER PLATFORM", "spc": 60}, font=HEAD_SB, size=12, color=WHITE)
add_text(s, X(0), 2.5, SPAN(5), 2.3, "Een CRM bouwen als functioneel consultant, met Claude",
         font=HEAD, size=38, color=WHITE, line_spacing=1.0)
add_text(s, X(0), 4.9, SPAN(4), 1.0,
         "Een sales-app op Power Platform. Wat werkte, wat niet, en hoe je het zelf probeert.",
         font=HEAD, size=18, color=WHITE)
add_text(s, X(0), SH - M - 0.3, SPAN(3), 0.3, "Bram van der Weijden", size=13, color=WHITE)
sh_h = 1.7
slant(s, panel_w - M - sh_h * 196 / 269, SH - M - sh_h, sh_h, white=True)
picture(s, IMG / "04-opportunity-process.png", panel_w + 0.55, 1.6, 6.4, 4.3, halign="left", valign="middle")
notes(s, "Ik laat vandaag zien hoe ik in een paar weken een werkende sales-app heb gebouwd die lijkt op wat wij "
         "dagelijks implementeren in D365 Sales. Niet omdat het een product wordt, maar om uit te zoeken wat een "
         "functioneel consultant tegenwoordig zelf kan. Het gaat vooral over de werkwijze; de app is het bewijs.",
      "1 min")

# 2. Agenda --------------------------------------------------------------------
s = content_slide("Vandaag", "Wat je vandaag ziet")
add_text(s, X(0), TOP, SPAN(3), 2.0, "Het gaat vooral over de werkwijze. De app is het bewijs.",
         font=HEAD, size=22, line_spacing=1.1)
agenda = [("Opening en de app in één blik", "2"), ("Werkwijze en skills", "5"),
          ("Het proces: van kaart tot bouwticket", "7"), ("Kwaliteit en tests", "2"),
          ("De app: techniek, v1, v2 en v3", "4"), ("Wat dit betekent voor een consultant", "2,5"),
          ("Demo", "10"), ("Afsluiting en vragen", "3")]
row = 0.56
for i, (t, mins) in enumerate(agenda):
    yy = TOP + i * row
    line(s, X(4), yy, RIGHT, yy, color=DARK if i == 0 else RULE)
    add_text(s, X(4), yy + 0.14, SPAN(4), 0.4, t, size=16)
    add_text(s, X(8), yy + 0.14, SPAN(1), 0.4, f"{mins} min", size=14, color=GREY, align=PP_ALIGN.RIGHT)
notes(s, "Kort de opbouw: eerst hoe ik werk, dan het proces, dan de app zelf, en daarna de demo. "
         "Doel: onthoud één zin. Als functioneel consultant kun je een echte, geteste bedrijfsapp op Power Platform "
         "bouwen door zelf de beslissingen te nemen, terwijl Claude de code schrijft.", "0,5 min")

# 3. Waarom dit project ---------------------------------------------------------
s = content_slide("Aanleiding", "Waarom dit project")
lined_list(s, X(0), TOP, SPAN(4), [
    ("Een alternatief voor D365 Sales", "Eerst voor mijn eigen team of één klant"),
    ("Lagere licentiekosten, eenvoudigere UI", None),
    ("Oefenproject met professionele standaarden", "'Oefenen' mag de lat niet verlagen"),
    ("Front-end: Power Apps code app", None),
    ("Data: Dataverse", None)], 0.86)
px, py, pw, ph = picture(s, IMG / "02-pipeline.png", X(4) + 0.2, TOP, SPAN(5) - 0.2, 4.3, halign="right")
caption(s, px, py + ph + 0.1, pw, "Het Pipeline-bord in de app. Alle data is fictief.")
notes(s, "D365 Sales is breed, maar voor een klein salesteam vaak zwaar en duur. Ik wilde weten hoe ver je komt met "
         "alleen Dataverse en een eigen front-end. Het is een oefenproject, maar ik behandel het alsof het ooit serieus "
         "wordt: nette architectuur, tests, vaste taal. \"Oefenen\" mag de lat niet verlagen.", "1 min")

# 4. Werkwijze (diagram A, opnieuw opgebouwd) ------------------------------------
s = content_slide("Werkwijze", "Telefoon, Claude Project en een PC met Claude Code")
bw, bh, gap = SPAN(3) - 0.7, 1.3, 0.7
xs = [X(0), X(3), X(6)]
row1 = [("Telefoon", "Claude-app: beslissingen nemen en de voortgang volgen", True),
        ("Claude Project", "Een thread per taak, met geheugen voor beslissingen", False),
        ("PC met Claude Code", "git, pac, pa en de Dataverse CLI, al ingelogd", False)]
for x, (t, d, hl) in zip(xs, row1):
    node(s, x, TOP, bw, bh, t, d, highlight=hl, title_size=16, sub_size=12)
for i in range(2):
    arrow(s, xs[i] + bw + 0.12, TOP + bh / 2, xs[i + 1] - 0.12, TOP + bh / 2)
bus_y, row2_y, row2_h = TOP + bh + 0.38, TOP + bh + 0.78, 1.1
gh_cx, dv_cx = xs[1] + bw / 2, xs[2] + bw / 2
line(s, dv_cx, TOP + bh, dv_cx, bus_y, color=GREY, width=1.0)
line(s, gh_cx, bus_y, dv_cx, bus_y, color=GREY, width=1.0)
arrow(s, gh_cx, bus_y, gh_cx, row2_y - 0.06)
arrow(s, dv_cx, bus_y, dv_cx, row2_y - 0.06)
node(s, xs[1], row2_y, bw, row2_h, "GitHub", "Code, Issues, pull requests en CI", title_size=16, sub_size=12)
node(s, xs[2], row2_y, bw, row2_h, "Dataverse", "dev-salesapp: tabellen, rollen, plug-ins en de code app",
     title_size=16, sub_size=12)
add_text(s, X(0), row2_y + 0.05, SPAN(3) - 0.5, 1.2,
         [{"text": "Ik", "font": HEAD_SB, "size": 14, "color": ACCENT},
          {"text": "Vragen beantwoorden, kiezen, reviewen. Een melding als ik iets moet beslissen.",
           "size": 12.5, "color": GREY, "space_before": 3}])
line(s, X(0), 5.55, RIGHT, 5.55)
add_text(s, X(0), 5.7, SPAN(9), 0.5, "Ik beslis en review. Claude bouwt.", font=HEAD, size=24)
notes(s, "Het grootste deel doe ik vanaf mijn telefoon. In een Claude Project start ik per taak een thread: een "
         "onderzoek, een uitvraag of een bouwticket. Die threads onthouden beslissingen. Voor werk dat op mijn eigen "
         "omgeving moet draaien, zoals solutions importeren of deployen, gebruikt Claude een PC waarop Claude Code "
         "staat, met git, pac, pa en de Dataverse CLI al ingelogd. Ik zit dus niet de hele dag achter een scherm; ik "
         "krijg een melding als ik iets moet beslissen.", "2 min")

# 5. Skills (diagram E, opnieuw opgebouwd) --------------------------------------
s = content_slide("Werkwijze", "Skills: herbruikbare werkwijzen")
cols = [("Power Platform skills", "Van Microsoft",
         ["Code apps: maken, koppelen, deployen", "Power Automate: flows bouwen en debuggen",
          "Model-driven apps, Power Pages en canvas apps"]),
        ("Dataverse skills", "Van Microsoft",
         ["Verbinden", "Metadata: tabellen en kolommen", "Data en queries", "Security-rollen", "Solutions"]),
        ("Matt Pocock skills", "Over de werkwijze",
         ["Wayfinder: kaart van beslissingen", "Grilling: vragen met aanbeveling", "Prototype: drie UI-varianten",
          "Research en domain modeling", "TDD, code review en pull requests"])]
for i, (t, tag, items) in enumerate(cols):
    x, w = X(i * 3), SPAN(3)
    line(s, x, TOP, x + w, TOP)
    add_text(s, x, TOP + 0.15, w, 0.4, t, font=HEAD_SB, size=18)
    add_text(s, x, TOP + 0.55, w, 0.3, tag.upper(), font=HEAD_SB, size=10, color=ACCENT)
    for j, it in enumerate(items):
        yy = TOP + 0.98 + j * 0.5
        line(s, x, yy, x + w, yy, color=RULE)
        add_text(s, x, yy + 0.1, w, 0.4, it, size=13.5)
line(s, X(0), 5.55, RIGHT, 5.55)
add_text(s, X(0), 5.7, SPAN(9), 0.5, "Een skill is een beschreven werkwijze voor Claude: zo doe je dit.",
         font=HEAD, size=22)
notes(s, "Een skill is een instructie-pakket: \"zo doe je dit\". De Power Platform- en Dataverse-skills komen van "
         "Microsoft en kennen de CLI's en patronen, bijvoorbeeld hoe je een tabel aan een code app koppelt. De Matt "
         "Pocock-skills gaan over werkwijze: eerst vragen stellen, dan een plan, dan bouwen met tests. Die zijn niet "
         "specifiek voor software; je kunt ze ook voor andere beslissingen gebruiken.", "2 min")

# 5b. Matt Pocock: de filosofie (bron: github.com/mattpocock/skills) -------------
s = content_slide("Werkwijze", "Matt Pocock: echte engineering, geen vibe coding")
add_text(s, X(0), TOP, SPAN(3) - 0.2, 1.6, "De basis van goed software bouwen telt meer dan ooit.",
         font=HEAD, size=21, line_spacing=1.1)
add_text(s, X(0), TOP + 1.7, SPAN(3) - 0.2, 2.6,
         [{"text": "Matt Pocock is bekend van Total TypeScript en schrijft nu over AI-ontwikkeling (AI Hero). Zijn "
                   "skills zijn klein, aan te passen en werken met elk model.", "size": 12.5, "color": DARK},
          {"text": "Frameworks als Spec-Kit of BMAD nemen het proces over. Zijn skills niet: jij houdt de regie.",
           "size": 12.5, "color": DARK, "space_before": 8}])
caption(s, X(0), BOTTOM - 0.3, SPAN(3), "Bron: github.com/mattpocock/skills")
tx = X(3)
cw = [SPAN(2) + 0.25, SPAN(2) + 0.1, SPAN(2) - 0.35]
cx = [tx, tx + cw[0] + G, tx + cw[0] + cw[1] + 2 * G]
for x, w, hd in zip(cx, cw, ["Het probleem", "De fix", "Bij ons"]):
    add_text(s, x, TOP, w, 0.3, hd.upper(), font=HEAD_SB, size=10, color=ACCENT)
pocock = [
    ("Doet niet wat ik wil", "\"No-one knows exactly what they want\"  ·  Pragmatic Programmer",
     "Grilling: de agent vraagt jou eerst uit", "Wayfinder-kaart en grilling, zoals issue #3"),
    ("De agent praat te veel", "Gedeelde taal  ·  Eric Evans, Domain-Driven Design",
     "Eén woordenlijst en vastgelegde keuzes", "GLOSSARY.md: \"Opportunity\", niet \"Deal\". ADR 0001-0003"),
    ("De code werkt niet", "\"The rate of feedback is your speed limit\"  ·  Pragmatic Programmer",
     "Snelle feedback: types, tests eerst (TDD), kleine stappen", "Ruim 40 testbestanden, CI bij elke pull request"),
    ("Een grote modderbal", "Kent Beck en John Ousterhout, A Philosophy of Software Design",
     "Elke dag aan het ontwerp werken: diepe modules, simpele interface", "SalesGateway: één interface tussen app en data")]
rh = 1.0
for i, (prob, src, fix, ons) in enumerate(pocock):
    yy = TOP + 0.4 + i * rh
    line(s, tx, yy, RIGHT, yy, color=DARK if i == 0 else RULE)
    add_text(s, cx[0], yy + 0.13, cw[0], rh - 0.15,
             [{"text": prob, "font": HEAD_SB, "size": 14},
              {"text": src, "size": 10, "color": GREY, "space_before": 3}])
    add_text(s, cx[1], yy + 0.15, cw[1], rh - 0.15, fix, size=12.5)
    add_text(s, cx[2], yy + 0.15, cw[2], rh - 0.15, ons, size=12.5, color=GREY)
notes(s, "Even over Matt Pocock zelf. Hij is bekend als TypeScript-docent en schrijft nu over werken met AI-agents. "
         "Zijn repo heet 'Skills for real engineers': echte engineering, geen vibe coding. Hij zet zich af tegen "
         "frameworks die het hele proces overnemen. Zijn skills zijn klein, je past ze zelf aan, en ze werken met elk "
         "model.\n\n"
         "Hij gaat uit van vier dingen die bij AI-agents steeds misgaan, en elk antwoord komt uit een klassiek boek.\n"
         "1. De agent doet niet wat je wilt. Niemand weet precies wat hij wil, dus laat de agent jou uitvragen: grilling.\n"
         "2. De agent praat te veel. Geef hem een vaste woordenlijst, zoals in Domain-Driven Design. Bij ons: "
         "GLOSSARY.md.\n"
         "3. De code werkt niet. Feedback is je snelheidslimiet: tests eerst, kleine stappen.\n"
         "4. Een grote modderbal. Agents maken code sneller, en dus ook sneller rommelig. Werk elke dag aan het ontwerp. "
         "Bij ons houdt de SalesGateway de app en de data uit elkaar.\n\n"
         "Voor ons als consultants is punt 1 het herkenbaarst: een klant weet ook nooit precies wat hij wil. Dat is "
         "gewoon uitvraag, en daar zijn we goed in.", "1 min")

# 6. Wayfinder -----------------------------------------------------------------
s = content_slide("Het proces · stap 1", "Wayfinder: een kaart van beslissingen")
lined_list(s, X(0), TOP, SPAN(4), [
    ("Begint met het doel", "\"Een bouwbare spec voor v1\""),
    ("Maakt een kaart in GitHub", "Een issue met beslispunten als tickets"),
    ("Elk ticket is onderzoek, uitvraag of prototype", None),
    ("Voortgang zichtbaar", "Welke vragen zijn nog open?")], 0.95)
add_text(s, X(0), TOP + 4 * 0.95 + 0.15, SPAN(4), 0.6,
         "Claude doet voorstellen. Ik besluit.", font=HEAD, size=18)
px0, pw0 = X(4) + 0.2, SPAN(5) - 0.2
py0, ph0 = TOP, BOTTOM - TOP
box(s, px0, py0, pw0, ph0, outline=None, fill=LIGHT)
pad = 0.25
add_text(s, px0 + pad, py0 + 0.18, pw0 - 2 * pad, 0.35,
         {"runs": [("#2  ", {"color": GREY}), ("Map: v1 sales app spec", {})]}, font=HEAD_SB, size=14)
add_text(s, px0 + pad, py0 + 0.5, pw0 - 2 * pad, 0.3, "14 VAN 14 BESLISPUNTEN GESLOTEN", font=HEAD_SB, size=9.5,
         color=ACCENT)
tickets = [(3, "Uitvraag", "How a Lead moves from new to qualified or disqualified"),
           (4, "Uitvraag", "Opportunity stages and how a deal closes"),
           (5, "Uitvraag", "What Qualify creates when the Account or Contact already exists"),
           (6, "Uitvraag", "Fields on Lead and Opportunity in v1"),
           (7, "Prototype", "How the pipeline board should look and behave"),
           (8, "Onderzoek", "Dataverse security setup for read-all, owner-and-manager-edit"),
           (9, "Onderzoek", "What Power Apps code apps can do with Dataverse today"),
           (10, "Onderzoek", "Which Dataverse activity tables fit notes and simple activities"),
           (12, "Taak", "Can the code app write polymorphic lookups: note, task, owner"),
           (13, "Taak", "Does manager hierarchy let a manager edit their report's Lead"),
           (18, "Uitvraag", "Which screens the app has and how a user moves between them"),
           (19, "Uitvraag", "How the app is packaged and promoted from dev to production"),
           (20, "Uitvraag", "What sample data v1 ships with for demos and testing"),
           (21, "Uitvraag", "Who can reassign a record, and what the sales role is called")]
ty0 = py0 + 0.9
th = (py0 + ph0 - 0.15 - ty0) / len(tickets)
for i, (n, kind, title) in enumerate(tickets):
    yy = ty0 + i * th
    line(s, px0 + pad, yy, px0 + pw0 - pad, yy, color=DARK if i == 0 else RULE, width=0.5)
    add_text(s, px0 + pad, yy + 0.06, 0.95, th, kind.upper(), font=HEAD_MD, size=7.5, color=GREY)
    add_text(s, px0 + pad + 0.95, yy + 0.05, pw0 - 2 * pad - 0.95, th,
             {"runs": [(f"#{n}  ", {"color": GREY}), (title, {})]}, size=10)
notes(s, "Ik begon niet met bouwen maar met een kaart. Wayfinder zet het doel neer en splitst het op in beslispunten: "
         "wat is een Lead-status, wat gebeurt er bij Qualify, wie mag wat? Elk punt is een GitHub-ticket. Dat dwingt me "
         "om eerst te bedenken wat ik wil, in plaats van Claude maar iets te laten verzinnen. Voor v1 waren dat 14 "
         "beslispunten, rechts zie je ze allemaal. Let op: Claude doet voorstellen, ik besluit.\n\n"
         "Live laten zien: github.com/Bramvdw94/power-platform-sales-app/issues/2", "1,5 min")

# 7. Grilling ------------------------------------------------------------------
s = content_slide("Het proces · stap 2", "Grilling: vragen met een aanbeveling")
lined_list(s, X(0), TOP, SPAN(4), [
    ("Claude stelt genummerde vragen", None),
    ("Bij elke vraag staat een aanbeveling", None),
    ("Ik antwoord \"ja\" of pas aan", None),
    ("Voorbeeld: de levenscyclus van een Lead", "Issue #3")], 0.8)
tx, tw = X(4) + 0.2, SPAN(5) - 0.2
add_text(s, tx, TOP - 0.02, tw, 0.4, "Levenscyclus van een Lead (#3)", font=HEAD_SB, size=15)
hy = TOP + 0.5
line(s, tx, hy, tx + tw, hy)
add_text(s, tx, hy + 0.1, 1.9, 0.3, "STATUS", font=HEAD_SB, size=9.5, color=GREY)
add_text(s, tx + 2.0, hy + 0.1, tw - 2.0, 0.3, "BETEKENIS", font=HEAD_SB, size=9.5, color=GREY)
statuses = [("Not Contacted", "Aangemaakt, nog niemand benaderd"),
            ("Working", "Owner is in gesprek en beoordeelt de fit"),
            ("Qualified", "Qualify maakte een Opportunity. Definitief"),
            ("Disqualified", "Afgesloten zonder deal. Kan heropend worden")]
for i, (st, mean) in enumerate(statuses):
    yy = hy + 0.45 + i * 0.55
    line(s, tx, yy, tx + tw, yy, color=RULE)
    add_text(s, tx, yy + 0.14, 1.9, 0.4, st, font=HEAD_SB, size=13)
    add_text(s, tx + 2.0, yy + 0.14, tw - 2.0, 0.4, mean, size=13)
ry = hy + 0.45 + 4 * 0.55
line(s, tx, ry, tx + tw, ry, color=RULE)
add_text(s, tx, ry + 0.3, tw, 0.3, "REDEN BIJ DISQUALIFY, UIT EEN VASTE LIJST", font=HEAD_SB, size=9.5, color=GREY)
cx = tx
cy = ry + 0.65
for reason in ["No budget", "No need", "Can't reach them", "Chose a competitor", "Duplicate", "Other"]:
    cw = 0.3 + len(reason) * 0.085
    if cx + cw > tx + tw:
        cx, cy = tx, cy + 0.45
    box(s, cx, cy, cw, 0.34, outline=DARK, width=0.75)
    add_text(s, cx, cy + 0.07, cw, 0.25, reason, size=10.5, align=PP_ALIGN.CENTER)
    cx += cw + 0.12
notes(s, "Deze sessie is eigenlijk net als een goed intakegesprek. Claude vraagt: welke statussen, welke overgangen, "
         "wat leg je vast bij disqualify, mag je heropenen? Bij elke vraag geeft het een aanbeveling, in dit geval "
         "gebaseerd op de standaard-statussen van Salesforce. Ik hoefde vaak alleen \"ja\" te zeggen of een detail aan "
         "te passen. Zo zijn beslissingen snel genomen en blijven ze vastgelegd in het ticket. Dat is precies wat wij "
         "als consultants ook doen, alleen nu met een gesprekspartner die alles uitwerkt.", "1,5 min")

# 8. Prototype -----------------------------------------------------------------
s = content_slide("Het proces · stap 3", "Prototype: drie varianten om uit te kiezen")
end = lined_list(s, X(0), TOP, SPAN(4), [
    ("Voor schermen eerst een prototype", "Pas daarna bouwen"),
    ("Meestal drie varianten: A, B en C", None),
    ("Ik kies, of combineer", "Bij Forecast werd het variant D"),
    ("Prototypes worden nooit gemerged", None)], 0.9)
caption(s, X(0), end + 0.15, SPAN(4), "Op de schermen: de gekozen Forecast-pagina (variant D) en het Team-tabblad.")
iw = 4.7
ih = iw * 900 / 1440
picture(s, IMG / "05-forecast.png", X(4) + 0.2, TOP, iw, ih, halign="left")
picture(s, IMG / "06-forecast-team.png", RIGHT - iw, BOTTOM - ih, iw, ih, halign="left")
notes(s, "Bij schermen is praten niet genoeg. Claude bouwt dan een prototype in de echte app, in het geheugen, met drie "
         "varianten naast elkaar. Voor de Forecast-pagina vond ik alle drie mooi en koos ik een combinatie. Ook voor de "
         "offerte-PDF bouwde Claude een prototype, en daar bedacht ik zelf een sjabloon-builder met drag-and-drop bij. "
         "Belangrijk: elk prototype wordt ook gecontroleerd in lichte en donkere modus, want formuliervelden waren in "
         "donkere modus in het begin onleesbaar.", "1,5 min")

# 8b. Licht en donker (optioneel) ----------------------------------------------
s = content_slide("Het proces · stap 3", "Elk prototype in licht én donker gecontroleerd")
half = (SPAN(9) - G * 2) / 2
for i, (f, cap_text) in enumerate([("04-opportunity-process.png", "Lichte modus"),
                                   ("12-dark-opportunity.png", "Donkere modus")]):
    x = X(0) + i * (half + G * 2)
    px, py, pw, ph = picture(s, IMG / f, x, TOP, half, 3.9, halign="left")
    caption(s, px, py + ph + 0.1, pw, cap_text)
add_text(s, X(0), 6.1, SPAN(9), 0.4,
         "Feedback als \"controleer dit voortaan altijd\" gaat in het projectgeheugen. Elke volgende thread neemt het mee.",
         size=14)
notes(s, "Optioneel, bij weinig tijd overslaan. Dit is dezelfde pagina in licht en donker. Dit soort feedback "
         "(\"controleer dit voortaan altijd\") sla ik op in het projectgeheugen, zodat elke volgende thread het "
         "meeneemt.", "0,5 min (optioneel)")

# 9. Spec en issues: de cirkel (diagram B, opnieuw opgebouwd) -------------------
s = content_slide("Het proces · stap 4", "Spec en issues: het hele proces is een cirkel")
lined_list(s, X(0), TOP, SPAN(3), [
    ("Per versie opnieuw", "v1, v2 en v3"),
    ("Spec = een GitHub issue", "Probleem, oplossing en user stories"),
    ("Bouwtickets zijn sub-issues", "Met afhankelijkheden"),
    ("Volgorde via \"blocked by\"", "Geen aparte planning nodig")], 1.0, size=15)
dx, dw = X(3) + 0.1, SPAN(6) - 0.1
nw = (dw - 3 * 0.42) / 4
nh, rows = 0.82, [TOP, TOP + 1.2, TOP + 2.4, TOP + 3.6]
cxs = [dx + i * (nw + 0.42) for i in range(4)]
loop = [(0, 0, "Wayfinder", "kaart en beslispunten", False), (1, 0, "Research", "uitzoeken", False),
        (2, 0, "Grilling", "vragen en aanbeveling", False), (3, 0, "Prototype", "drie varianten", False),
        (3, 1, "Spec", "user stories", True), (3, 2, "Issues", "bouwtickets", True),
        (3, 3, "Bouwen", "tests eerst (TDD)", False), (2, 3, "Deploy", "CI groen", False),
        (1, 3, "Demo", "feedback: wat nu?", False), (0, 3, "Nieuwe versie", "v1, v2, v3", False)]
for c, r, t, sub, hl in loop:
    node(s, cxs[c], rows[r], nw, nh, t, sub, highlight=hl, title_size=12, sub_size=10, align=PP_ALIGN.CENTER)
for (c1, r1, *_), (c2, r2, *_) in zip(loop, loop[1:]):
    if r1 == r2:
        y = rows[r1] + nh / 2
        if c2 > c1:
            arrow(s, cxs[c1] + nw + 0.06, y, cxs[c2] - 0.06, y)
        else:
            arrow(s, cxs[c1] - 0.06, y, cxs[c2] + nw + 0.06, y)
    else:
        x = cxs[c1] + nw / 2
        arrow(s, x, rows[r1] + nh + 0.06, x, rows[r2] - 0.06)
arrow(s, cxs[0] + nw / 2, rows[3] - 0.06, cxs[0] + nw / 2, rows[0] + nh + 0.06)
add_text(s, cxs[1], rows[1] + 0.2, cxs[2] + nw - cxs[1], 1.8,
         [{"text": "Per versie één keer rond", "font": HEAD, "size": 20, "align": PP_ALIGN.CENTER},
          {"text": "v1: 12 bouwtickets  ·  v2: 16  ·  v3: 32 klaar", "size": 12, "color": GREY,
           "align": PP_ALIGN.CENTER, "space_before": 6}], anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Als alle beslispunten dicht zijn, schrijft Claude een spec: een issue met het probleem, de oplossing en user "
         "stories. Daaruit komen bouwtickets. Voor v1 waren dat 12 tickets, voor v2 16, en voor v3 staan er 32 klaar. "
         "Elk ticket krijgt zijn eigen thread, wordt gebouwd, getest en gedeployed. Daarna begint de cirkel opnieuw met "
         "een nieuwe kaart voor de volgende versie. De volgorde van tickets staat vast door \"blocked by\"-regels, dus "
         "ik hoef geen planning bij te houden.", "1,5 min")

# 9a. Voorbeeld: zo hangen de issues aan elkaar (v1, uit GitHub) ----------------
s = content_slide("Het proces · voorbeeld", "Van kaart tot code: één spoor door GitHub")
gap = 0.42
cw5 = (SPAN(9) - 4 * gap) / 5
cx5 = [X(0) + i * (cw5 + gap) for i in range(5)]
steps = [("1 · Kaart", "/wayfinder"), ("2 · Beslispunten", "/grill-with-docs · /prototype · /research"),
         ("3 · Spec", "/to-spec"), ("4 · Bouwtickets", "/to-tickets"), ("5 · Code", "Bouwen, tests, CI")]
for x, (st, sk) in zip(cx5, steps):
    add_text(s, x, TOP, cw5, 0.3, st.upper(), font=HEAD_SB, size=10, color=ACCENT)
    add_text(s, x, TOP + 0.3, cw5, 0.3, sk, size=10.5, color=GREY)
cy, ch = TOP + 0.8, 2.4
node(s, cx5[0], cy, cw5, ch, "#2  Map: v1 sales app spec",
     "Doel: een bouwbare spec voor v1. Alle open vragen staan als sub-issue onder de kaart.", title_size=13, sub_size=11)
box(s, cx5[1], cy, cw5, ch)
dec = [("#3", "grilling", "Status van een Lead", True), ("#7", "prototype", "Het pipelinebord", False),
       ("#8", "research", "Dataverse-security", False), ("#12", "task", "Check in de omgeving", False)]
rh5 = (ch - 0.45) / 4
for j, (nr, lbl, t, hl) in enumerate(dec):
    yy = cy + 0.14 + j * rh5
    if j:
        line(s, cx5[1] + 0.16, yy, cx5[1] + cw5 - 0.16, yy, color=RULE)
    add_text(s, cx5[1] + 0.16, yy + 0.04, cw5 - 0.32, rh5 - 0.04,
             [{"runs": [(f"{nr} ", {"font": HEAD_SB, "color": ACCENT if hl else DARK}),
                        (lbl, {"color": GREY})], "size": 11},
              {"text": t, "size": 10.5, "color": DARK if hl else GREY}])
add_text(s, cx5[1] + 0.16, cy + ch - 0.34, cw5 - 0.32, 0.25, "+ 10 andere, 14 in totaal", size=10, color=GREY)
node(s, cx5[2], cy, cw5, ch, "#23  Spec: v1 sales app",
     "Probleem, oplossing en user stories, geschreven uit de 14 gesloten beslispunten.", title_size=13, sub_size=11)
node(s, cx5[3], cy, cw5, ch, "#31  Build: Qualify",
     "Eén van de 12 bouwtickets (#24 t/m #35). Met wat te bouwen en acceptatiecriteria.", title_size=13, sub_size=11)
node(s, cx5[4], cy, cw5, ch, "PR #49  Qualify",
     "Dialoog, Dataverse-gateway en Lead-actie. Tests groen, gemerged.", highlight=True, title_size=13, sub_size=11)
for i in range(4):
    arrow(s, cx5[i] + cw5 + 0.06, cy + ch / 2, cx5[i + 1] - 0.06, cy + ch / 2)
ly = cy + ch + 0.22
line(s, X(0), ly, RIGHT, ly)
add_text(s, X(0), ly + 0.08, SPAN(9), 0.3, "HOE HET AAN ELKAAR VASTZIT", font=HEAD_SB, size=10, color=GREY)
links = ["Elk beslispunt is een sub-issue van #2",
         "Besluit als comment in #3, woordenlijst bijgewerkt",
         "De spec verwijst terug: \"from the Lead lifecycle decision #3\"",
         "\"Part of #23\": bouwticket is sub-issue van de spec",
         "De PR sluit #31 bij het mergen"]
for x, t in zip(cx5, links):
    add_text(s, x, ly + 0.36, cw5, 0.9, t, size=11.5)
caption(s, X(0), BOTTOM - 0.2, SPAN(9),
        "Voorbeeld uit v1, Bramvdw94/power-platform-sales-app. Skillnamen uit github.com/mattpocock/skills.")
notes(s, "Laat ik één spoor volgen, zodat je ziet hoe het in GitHub aan elkaar hangt.\n"
         "1. Wayfinder maakt kaart #2: het doel en alle open vragen, elk als sub-issue.\n"
         "2. Die vragen hebben een soort: grilling (ik beslis), prototype (ik kies een scherm), research (Claude zoekt "
         "uit) of een korte taak in de omgeving. Bij #3 vroeg Claude hoe een Lead van nieuw naar gekwalificeerd gaat. "
         "Mijn besluit staat als comment in het issue, en de woordenlijst is bijgewerkt.\n"
         "3. Als alle 14 dicht zijn, schrijft /to-spec issue #23. Daarin staat letterlijk 'from the Lead lifecycle "
         "decision #3'. Elke regel is dus terug te voeren op een besluit.\n"
         "4. /to-tickets knipt de spec in bouwtickets, als sub-issues van de spec. #31 is Qualify.\n"
         "5. Claude bouwt #31 met tests eerst; pull request #49 sluit het ticket.\n"
         "Als een klant vraagt 'waarom werkt Qualify zo?', klik ik van de code terug naar mijn besluit.", "1 min")

# 9b. In cijfers (uit GitHub Issues) -------------------------------------------
s = content_slide("Het proces in cijfers", "Drie keer de cirkel rond")
versions = [("v1", "Gebouwd en gedeployed", "14", "12", "Kaart #2  ·  spec #23"),
            ("v2", "Gebouwd en gedeployed", "19", "16", "Kaart #58  ·  spec #91"),
            ("v3", "Spec klaar, bouw gepland", "24", "32", "Kaart #131  ·  spec #168")]
for i, (v, status, dec, builds, refs) in enumerate(versions):
    x, w = X(i * 3), SPAN(3)
    line(s, x, TOP, x + w, TOP)
    add_text(s, x, TOP + 0.12, w, 0.5, v, font=HEAD, size=28)
    add_text(s, x, TOP + 0.68, w, 0.3, status.upper(), font=HEAD_SB, size=10, color=ACCENT if i == 2 else GREY)
    for j, (num, lbl) in enumerate([(dec, "beslispunten op de kaart"), (builds, "bouwtickets")]):
        yy = TOP + 1.12 + j * 1.38
        line(s, x, yy, x + w, yy, color=RULE)
        add_text(s, x, yy + 0.1, w, 0.8, num, font=HEAD, size=44, color=BLACK, line_spacing=1.0)
        add_text(s, x, yy + 0.92, w, 0.3, lbl, size=13, color=GREY)
    yy = TOP + 1.12 + 2 * 1.38
    line(s, x, yy, x + w, yy, color=RULE)
    add_text(s, x, yy + 0.12, w, 0.3, refs, size=12)
caption(s, X(0), BOTTOM - 0.2, SPAN(9),
        "Bron: GitHub Issues in Bramvdw94/power-platform-sales-app (125 issues), oktober 2026. "
        "Beslispunten = onderzoek, uitvraag, prototype en taak.")
notes(s, "Zo ziet het papieren spoor eruit. Elke versie begint met een kaart vol beslispunten en eindigt in "
         "bouwtickets. v1 had 14 beslispunten en 12 bouwtickets, v2 19 en 16, en voor v3 zijn 24 beslispunten "
         "gesloten en staan 32 bouwtickets klaar. Alles is terug te vinden in GitHub.", "0,5 min")

# 10. Het Matt Pocock-proces ----------------------------------------------------
s = content_slide("Werkwijze", "Het Matt Pocock-proces in het kort")
lined_list(s, X(0), TOP, SPAN(5), [
    ("Eerst begrijpen", "Research en grilling"),
    ("Dan zien", "Prototype"),
    ("Dan vastleggen", "Spec en issues"),
    ("Dan bouwen", "Tests eerst (TDD), kleine stappen"),
    ("Vaste taal", "GLOSSARY.md en ADR's")], 0.88, numbered=True, title_font=HEAD_SB, size=16)
px0, pw0 = X(5) + 0.2, SPAN(4) - 0.2
box(s, px0, TOP, pw0, BOTTOM - TOP, outline=None, fill=LIGHT)
add_text(s, px0 + 0.25, TOP + 0.2, pw0 - 0.5, 0.4, "Vastgelegd in de repo", font=HEAD_SB, size=15)
lined_list(s, px0 + 0.25, TOP + 0.75, pw0 - 0.5, [
    ("GLOSSARY.md", "\"Opportunity\", niet \"Deal\" of \"Sale\""),
    ("ADR 0001", "Process engine: plug-in plus app-evaluatie"),
    ("ADR 0002", "Decision rules roepen Jev aan na het opslaan"),
    ("ADR 0003", "Approvals in een eigen tabel")], 0.88, size=14, title_font=BODY_MD)
notes(s, "De kern: neem beslissingen vóór je bouwt, leg ze vast, en bouw in kleine, geteste stappen. De glossary is mijn "
         "favoriet: \"Opportunity\", niet \"Deal\" of \"Sale\". Elke thread leest hem, dus Claude gebruikt overal "
         "dezelfde woorden. Grote keuzes, zoals hoe de proces-engine draait, staan als ADR in de repo, met het waarom "
         "erbij.", "1 min")

# 11. Kwaliteit (diagram D, opnieuw opgebouwd) -----------------------------------
s = content_slide("Kwaliteit", "Codekwaliteit en tests, een gevolg van dit proces")
a_x, a_w = X(0), 3.3
b_x, b_w = a_x + a_w + 0.6, 3.3
c_x = b_x + b_w + 0.75
c_w = RIGHT - c_x
r_y, r_h = TOP + 0.45, 1.1
node(s, a_x, r_y, a_w, r_h, "Code app", "React en TypeScript. Pagina's getest via de gateway")
node(s, b_x, r_y, b_w, r_h, "SalesGateway", "Eén interface: de enige naad tussen app en data", highlight=True)
c1_y, c2_y, c_h = TOP, TOP + 1.15, 0.95
node(s, c_x, c1_y, c_w, c_h, "In-memory twin", "Snelle tests, en een demo-modus zonder omgeving")
node(s, c_x, c2_y, c_w, c_h, "Dataverse gateway", "De echte omgeving, met plug-ins")
arrow(s, a_x + a_w + 0.06, r_y + r_h / 2, b_x - 0.06, r_y + r_h / 2)
bus_x = b_x + b_w + 0.35
line(s, b_x + b_w, r_y + r_h / 2, bus_x, r_y + r_h / 2, color=GREY, width=1.0)
line(s, bus_x, c1_y + c_h / 2, bus_x, c2_y + c_h / 2, color=GREY, width=1.0)
arrow(s, bus_x, c1_y + c_h / 2, c_x - 0.06, c1_y + c_h / 2)
arrow(s, bus_x, c2_y + c_h / 2, c_x - 0.06, c2_y + c_h / 2)
qy = 4.35
for i, (t, d) in enumerate([("Pure regels", "In src/sales: fase, kans, close, forecast en offertetotalen"),
                            ("Gedeelde testcases", "Dezelfde JSON-cases voor de TypeScript- en de C#-engine"),
                            ("Rechten getest", "xUnit voor plug-ins, check-scripts als Megan, Adele en Diego")]):
    x, w = X(i * 3), SPAN(3)
    line(s, x, qy, x + w, qy)
    add_text(s, x, qy + 0.14, w, 1.0, [{"text": t, "font": HEAD_SB, "size": 15},
                                        {"text": d, "size": 13, "color": GREY, "space_before": 3}])
line(s, X(0), 5.75, RIGHT, 5.75, color=RULE)
add_text(s, X(0), 5.88, SPAN(9), 0.4,
         {"runs": [("CI bij elke pull request: ", {"font": BODY_MD}),
                   ("lint, typecheck en tests.  Ruim 40 testbestanden voor de app, plus een apart testproject voor de "
                    "plug-ins.", {})]}, size=13)
notes(s, "Omdat de spec vooraf beschrijft wat de app moet doen, kan Claude eerst een test schrijven en daarna de code. "
         "De app praat met data via één interface; in de tests zit daar een geheugenversie achter. Daardoor draait de "
         "hele app ook zonder omgeving, en dat gebruik ik straks voor de demovideo. In de repo staan ruim 40 "
         "testbestanden voor de app en een apart testproject voor de plug-ins. Voor beveiliging draaien check-scripts "
         "als Megan, Adele en Diego om te zien dat iedereen precies ziet wat hij mag zien.", "2 min")

# 12. De app en de techniek (diagram F, opnieuw opgebouwd) -----------------------
s = content_slide("De app", "De app en de techniek")
dw = SPAN(6)
nw = (dw - 2 * 0.5) / 3
xs = [X(0) + i * (nw + 0.5) for i in range(3)]
top_nodes = [("Front-end", "Power Apps code app: React, Vite, TypeScript, Tailwind"),
             ("Dataverse", "Account en Contact standaard, eigen tabellen met prefix vibe_"),
             ("Plug-in (C#)", "Eén generieke engine voor processen en business rules")]
bot_nodes = [("Desktop én telefoon", "Zijbalk of tabs onderin. Geen native app"),
             ("Beveiliging", "Rol plus manager-hiërarchie"),
             ("Solution", "Eén unmanaged solution: SalesApp")]
h1, h2, y2 = 1.35, 1.05, TOP + 1.35 + 0.45
for i in range(3):
    node(s, xs[i], TOP, nw, h1, *top_nodes[i], highlight=(i == 1), title_size=15, sub_size=11.5)
    node(s, xs[i], y2, nw, h2, *bot_nodes[i], title_size=13, sub_size=11)
    line(s, xs[i] + nw / 2, TOP + h1, xs[i] + nw / 2, y2, color=GREY, width=1.0)
for i in range(2):
    arrow(s, xs[i] + nw + 0.06, TOP + h1 / 2, xs[i + 1] - 0.06, TOP + h1 / 2)
ly = y2 + h2 + 0.45
line(s, X(0), ly, X(0) + dw, ly)
add_text(s, X(0), ly + 0.14, dw, 1.0,
         [{"text": "Let op: licenties", "font": HEAD_SB, "size": 15},
          {"text": "Elke gebruiker heeft Power Apps Premium of een per-app-licentie nodig. Niet gratis, wel zonder "
                   "D365 Sales-licentie.", "size": 13, "color": GREY, "space_before": 3}])
px, py, pw, ph = picture(s, IMG / "10-phone-pipeline.png", X(6) + 0.2, TOP, SPAN(3) - 0.2, BOTTOM - TOP - 0.4)
caption(s, px, py + ph + 0.1, pw, "De pipeline op een telefoon")
notes(s, "Een code app is een gewone webapp in je eigen code, met Power Platform voor login, connectors en hosting. Voor "
         "ons is het interessant omdat je de volledige UI-vrijheid hebt, maar nog steeds op Dataverse en "
         "Entra-beveiliging draait. Let op: elke gebruiker heeft Power Apps Premium of een per-app licentie nodig. Dat is "
         "dus niet gratis, maar wel zonder D365 Sales-licentie.", "1,5 min")

# 13. v1 en v2 (diagram C, deel 1) ----------------------------------------------
s = content_slide("De app", "Wat is er gebouwd: v1 en v2")
vcols = [("v1: de basis", "Gebouwd  ·  12 tickets  ·  spec #23",
          ["Accounts en Contacts", "Leads met Disqualify", "Opportunities en Qualify", "Pipeline-bord en Funnel",
           "Activiteiten en Home", "Beveiligingsrol en manager-hiërarchie", "Voorbeelddata"]),
         ("v2: de motor", "Gebouwd  ·  16 tickets  ·  spec #91",
          ["Proces-engine: stappen en checks", "Business rules builder", "Activiteiten-rollup op Account",
           "Producten en offertes (PDF)", "Dashboards en forecast", "Quota per verkoper"])]
for i, (t, tag, items) in enumerate(vcols):
    x, w = X(i * 3), SPAN(3)
    line(s, x, TOP, x + w, TOP)
    add_text(s, x, TOP + 0.14, w, 0.4, t, font=HEAD_SB, size=18)
    add_text(s, x, TOP + 0.55, w, 0.3, tag.upper(), font=HEAD_SB, size=9.5, color=ACCENT)
    for j, it in enumerate(items):
        yy = TOP + 0.95 + j * 0.5
        line(s, x, yy, x + w, yy, color=RULE)
        add_text(s, x, yy + 0.11, w, 0.4, it, size=13.5)
px, py, pw, ph = picture(s, IMG / "08-automation.png", X(6) + 0.2, TOP, SPAN(3) - 0.2, 3.0)
add_text(s, px, py + ph + 0.15, pw, 1.2,
         [{"text": "Automation: hier bouwt de manager processen en business rules.", "size": 12, "color": GREY},
          {"text": "Beide versies draaien in dev-salesapp, alles bijgehouden in GitHub Issues.", "size": 12,
           "color": GREY, "space_before": 6}])
notes(s, "v1 is de basis van elke sales-tool: leads, qualify, opportunities, een pipeline-bord en activiteiten, met één "
         "rol en een managerhiërarchie. v2 voegt de dingen toe waar D365 Sales sterk in is: een proces-engine met "
         "stappen en checks die managers zelf bouwen, business rules, producten met offertes en PDF-sjablonen, en een "
         "forecast met quota. Dat zie je straks in de demo.", "1,5 min")

# 14. v3 (diagram C, deel 2) ----------------------------------------------------
s = content_slide("Gepland", "Gepland: v3, met Jev en Mistral")
end = lined_list(s, X(0), TOP, SPAN(5), [
    ("Outlook en Teams", "Mail, afspraken en notificaties"),
    ("Gebruiksgemak", "Opgeslagen views, import, verwijderen, duplicaten"),
    ("Proces-engine ronde 2", "Approvals en Decision rules met Jev"),
    ("Forecast-diepgang, web lead capture", "Kortingen en Orders"),
    ("AI-assist met Mistral", "Samenvattingen en conceptmails"),
    ("Geen betaalde Azure", "De dev-tenant heeft geen krediet")], 0.7, size=15, desc_size=12)
caption(s, X(0), end + 0.1, SPAN(5), "Spec #168  ·  32 bouwtickets (#169 t/m #200)  ·  Test/Prod-pipeline bewust niet in v3")
px0, pw0 = X(5) + 0.2, SPAN(4) - 0.2
ph0 = (BOTTOM - TOP - 0.2) / 2
for i, (tag, cost, txt) in enumerate([
        ("Jev  ·  beslist  ·  per call", "ca. $0,00006", "\"Hoe heet is deze webshop-lead?\" Ca. 1.500 tokens à $0,042 "
                                                          "per miljoen. Een mens kijkt mee als het onzeker is."),
        ("Mistral  ·  schrijft", "Gratis", "Gratis API-key, zonder creditcard. Samenvatting, conceptmail, volgende "
                                          "stap. Alleen op verzoek en alleen met testdata.")]):
    y = TOP + i * (ph0 + 0.2)
    box(s, px0, y, pw0, ph0, outline=None, fill=LIGHT)
    add_text(s, px0 + 0.25, y + 0.2, pw0 - 0.5, 0.3, tag.upper(), font=HEAD_SB, size=10, color=ACCENT)
    add_text(s, px0 + 0.25, y + 0.5, pw0 - 0.5, 0.7, cost, font=HEAD, size=34, color=BLACK)
    add_text(s, px0 + 0.25, y + 1.25, pw0 - 0.5, ph0 - 1.35, txt, size=12.5, color=DARK)
notes(s, "v3 staat als spec klaar (issue #168), met 32 bouwtickets. Twee dingen wil ik uitlichten. Jev is een goedkoop "
         "model dat alleen beslissingen neemt: \"hoe heet is deze webshop-lead?\", met een vast antwoord. Het draait na "
         "het opslaan en kan acties starten, en een mens kijkt mee als het onzeker is. Mistral schrijft: samenvatting "
         "van een account, concept-opvolgmail, voorstel voor de volgende stap. Beide alleen op verzoek. Waarom deze "
         "twee? De dev-tenant heeft geen Azure-krediet en ik koppel geen creditcard, dus het moet gratis of zeer "
         "goedkoop zijn. Jev kost ongeveer 0,006 cent per call: 1.000 beslissingen per maand is zo'n 6 cent. Voor "
                  "Mistral is er een gratis API-key, zonder creditcard. Let op: gratis verzoeken mag Mistral gebruiken om te "
                  "trainen, dus alleen testdata. Een Test/Prod-pipeline staat bewust niet in v3.", "1,5 min")

# 15. Functioneel consultant ----------------------------------------------------
s = content_slide("Wat dit betekent", "Zo kan het ook als functioneel consultant")
add_text(s, X(0), TOP, SPAN(9), 0.6, "Ik schrijf geen code. Ik neem beslissingen.", font=HEAD, size=26)
ty = TOP + 0.95
heads = ["Ik", "Claude", "Samen"]
rows_t = [("Doel en scope", "Onderzoek en alternatieven", "Spec"),
          ("Kiezen uit varianten", "Prototypes en code", "Review"),
          ("Acceptatie", "Tests en deploy", "Demo en bijsturen")]
for i, hd in enumerate(heads):
    x, w = X(i * 3), SPAN(3)
    line(s, x, ty, x + w, ty)
    add_text(s, x, ty + 0.12, w, 0.4, hd, font=HEAD_SB, size=17, color=ACCENT if i == 0 else BLACK)
    for j, r in enumerate(rows_t):
        yy = ty + 0.6 + j * 0.55
        line(s, x, yy, x + w, yy, color=RULE)
        add_text(s, x, yy + 0.13, w, 0.4, r[i], size=15)
by = ty + 0.6 + 3 * 0.55 + 0.35
for i, t in enumerate(["Mijn gereedschap: vragen beantwoorden, schermen kiezen, specs lezen.",
                       "Eigen domeinkennis is de bottleneck, niet de techniek.",
                       "Review blijft mijn werk: klopt dit met hoe sales werkt?"]):
    x, w = X(i * 3), SPAN(3)
    line(s, x, by, x + w, by)
    add_text(s, x, by + 0.14, w, 0.8, t, size=13, color=GREY)
notes(s, "Het verrassende: de techniek is niet de bottleneck. Wat ik als consultant altijd al deed, zoals goede vragen "
         "stellen, scope bewaken en keuzes afwegen, is nu juist het werk. Ik controleer ook veel: bij elke deploy test "
         "ik met de demo-gebruikers.", "1 min")

# 16. Grenzen, kosten en risico's -------------------------------------------------
s = content_slide("Wat dit betekent", "Eerlijk: grenzen, kosten en risico's")
lined_list(s, X(0), TOP, SPAN(6), [
    ("Claude kan fout zitten", "Tests en review blijven nodig"),
    ("Alleen een dev-tenant", "Geen Test/Prod of pipeline, bewust voor deze fase"),
    ("Licenties", "Gebruikers hebben Power Apps Premium nodig"),
    ("AI-kosten", "Jev een fractie van een cent per call, Mistral gratis met limieten"),
    ("Data", "Geen echte klantdata in dit project")], 0.86, title_font=HEAD_SB, size=16)
qx, qw = X(6) + 0.2, SPAN(3) - 0.2
line(s, qx, TOP, qx + qw, TOP, color=ACCENT, width=2.0)
add_text(s, qx, TOP + 0.25, qw, 3.0,
         "Zonder duidelijke spec bouwt Claude iets plausibels maar onjuists. Met een spec en tests werkt het "
         "verrassend goed.", font=HEAD, size=20, line_spacing=1.15)
notes(s, "Dit is geen product. Het draait in één ontwikkelomgeving, zonder acceptatie en productie, en dat is een "
         "bewuste keuze voor deze fase. De app is niet bedoeld om klanten van D365 weg te halen. Wel heb ik in het "
         "proces veel geleerd over wat wel en niet werkt: zonder duidelijke spec bouwt Claude iets plausibels maar "
         "onjuists, met een spec en tests werkt het verrassend goed.", "1 min")

# 17. Overgang naar de demo -----------------------------------------------------
s = content_slide("Demo", "Nu de app zelf")
lined_list(s, X(0), TOP, SPAN(4), [
    ("Demo van ongeveer 10 minuten", None),
    ("Ingelogd als systeembeheerder", "In de live app op dev-salesapp"),
    ("Alle voorbeelddata is fictief", None),
    ("Plan B", "Dezelfde app in geheugenmodus, of de opgenomen video op de volgende dia")], 0.95)
picture(s, IMG / "01-home.png", X(4) + 0.2, TOP, SPAN(5) - 0.2, 4.4, halign="right")
notes(s, "Ik laat nu de app zien. Alle data is fictief. Als de verbinding hapert, heb ik een versie die zonder "
         "omgeving draait en een opgenomen video.\n\n"
         "Live demo in de app op dev-salesapp, ingelogd als systeembeheerder. Plan B: lokaal met npm run dev:memory.",
      "0,5 min")

# 18. Demovideo -----------------------------------------------------------------
s = content_slide("Demo", "De app in twee minuten")
vw = SPAN(6)
vh = vw * 9 / 16
mv = s.shapes.add_movie(str(VIDEO), Inches(X(0)), Inches(TOP - 0.1), Inches(vw), Inches(vh),
                        poster_frame_image=str(poster_frame()), mime_type="video/mp4")
mv.name = "Demovideo sales-app"
round_corners(mv, vw, vh)
caption(s, X(0), TOP - 0.1 + vh + 0.08, vw, "Opgenomen in geheugenmodus met fictieve data, ca. 1:50. Klik om af te spelen.")
sx, sw_ = X(6) + 0.2, SPAN(3) - 0.2
add_text(s, sx, TOP - 0.1, sw_, 0.35, "Wat je ziet", font=HEAD_SB, size=15)
scenes = ["Home: \"Mijn werk\"", "Pipeline: bord en Funnel", "Opportunity met procesbalk",
          "Business rule: waarschuwing", "Activiteiten op de deal", "Forecast: team en quota",
          "Producten en offertes", "Automation: processen en regels", "Nieuwe lead en Qualify", "Vooruitblik op v3"]
sy0 = TOP + 0.35
srow = (TOP - 0.1 + vh - sy0) / len(scenes)
for i, sc in enumerate(scenes):
    yy = sy0 + i * srow
    line(s, sx, yy, sx + sw_, yy, color=DARK if i == 0 else RULE, width=0.5)
    add_text(s, sx, yy + 0.07, sw_, srow, {"runs": [(f"{i + 1:02d}  ", {"font": HEAD_MD, "color": ACCENT}),
                                                     (sc, {})]}, size=11.5)
notes(s, "Speel de video af als plan B, of als intro van 2 minuten voordat je live gaat.\n\n"
         "Demoscript live (10 min):\n"
         "0-1  Home: \"My work\", taken, deals, KPI-tegels. \"Dit ziet een verkoper bij het openen.\"\n"
         "1-2  Pipeline: bord en Funnel, gewogen waarde, sleep een kaart. \"Net als in D365, maar lichter.\"\n"
         "2-4  Opportunity met proces: procesbalk, Key fields, \"To move on\", Guidance. \"De manager bepaalt per stap "
         "wat nodig is.\"\n"
         "4-5  Business rule: waarschuwing \"Next step is empty\", vul in, weg. \"Geen code: de manager bouwde deze "
         "regel zelf.\"\n"
         "5-6  Automation: tabs Processes en Business rules. \"Hier bouwt de manager ze.\"\n"
         "6-7  Offerte: producten, Quote templates, PDF-voorbeeld. \"Eigen sjablonen, met drag-and-drop-blokken.\"\n"
         "7-8  Forecast: maand/kwartaal, Team, Quotas. \"Commit, Best case en quota per verkoper.\"\n"
         "8-9  Beveiliging: wissel naar Adele (alleen eigen deals bewerken) en Diego (buiten het team). \"Iedereen "
         "leest, owner en manager bewerken.\"\n"
         "9-10 Teaser v3: prototype van de Decision rule editor of de e-mail/Teams UI. \"Dit is wat er nu aankomt.\"\n\n"
         "Verwijs ook naar GitHub: open 1 minuut issue #2 (kaart), #168 (spec v3) en een bouwticket.\n"
         "Als iets misgaat: \"Dat is precies waarom er tests zijn.\" Herlaad, of speel de video af.", "10 min")

# 19. Probeer het zelf ----------------------------------------------------------
s = content_slide("Probeer het zelf", "Wat nu? Probeer het zelf")
lined_list(s, X(0), TOP, SPAN(6), [
    ("Kies een klein probleem uit je eigen werk", None),
    ("Laat Claude je uitvragen", "Grilling, nog zonder te bouwen"),
    ("Laat een prototype maken en kies", None),
    ("Schrijf de spec, bouw per ticket", None),
    ("Controleer altijd zelf", "Met testgebruikers")], 0.86, numbered=True, size=16)
sl_h = 2.2
slant(s, RIGHT - sl_h * 196 / 269, TOP, sl_h)
add_text(s, X(6) + 0.2, TOP + sl_h + 0.35, SPAN(3) - 0.2, 2.0,
         [{"text": "Je hoeft geen code te kunnen, wel te weten wat je wilt.", "font": HEAD, "size": 18},
          {"runs": [("github.com/Bramvdw94/power-platform-sales-app", {"link": REPO_URL})],
           "size": 12, "color": ACCENT, "space_before": 12}])
notes(s, "Wil je dit proberen? Begin klein. Je hoeft geen code te kunnen, maar wel te weten wat je wilt. Ik help graag "
         "met een eerste grilling-sessie.", "2 min, daarna Q&A")

# 20. Vragen --------------------------------------------------------------------
s = content_slide("Vragen", "Vragen?")
qa = [("Wat kost dit?", "Claude-abonnement, Power Apps-licenties en bijna niets aan AI: Jev ca. $0,00006 per call, "
                        "Mistral via een gratis API-key. Geen Azure-abonnement."),
      ("Is dit veilig voor klantdata?", "Nu alleen fictieve data. Voor echte data: rechten in Dataverse, secrets in "
                                        "plug-in-configuratie, verwerkersovereenkomst voor AI."),
      ("Waarom geen D365 Sales?", "Dit is een experiment. D365 Sales heeft veel meer."),
      ("Kan een klant dit gebruiken?", "Mogelijk, maar zonder Test/Prod, ALM en support nog niet."),
      ("Hoeveel tijd kostte het?", "Niet gemeten. Geef hier je eigen inschatting.")]
for i, (q, a) in enumerate(qa):
    yy = TOP + i * 0.86
    line(s, X(0), yy, RIGHT, yy, color=DARK if i == 0 else RULE)
    add_text(s, X(0), yy + 0.15, SPAN(3), 0.6, q, font=HEAD_SB, size=15)
    add_text(s, X(3), yy + 0.15, SPAN(6), 0.7, a, size=13.5, color=GREY)
notes(s, "Verwachte vragen staan op de dia, met een kort antwoord. Bij \"Hoeveel tijd kostte het?\": geef je eigen "
         "inschatting, het is niet gemeten.")

prs.save(OUT)
print(f"Opgeslagen: {OUT} ({len(prs.slides)} dia's)")
